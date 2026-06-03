# Copyright (c) 2023-present Plane Software, Inc. and contributors
# SPDX-License-Identifier: AGPL-3.0-only
# See the LICENSE file for details.

"""Import/export system with pluggable formatters.

Defines :class:`BaseFormatter` and concrete CSV, JSON, and XLSX implementations
used by :class:`plane.utils.porters.exporter.DataExporter` to convert serialized
row dicts into file/string payloads for export and back into row dicts for
import. CSV output is sanitized through :mod:`plane.utils.csv_utils` to prevent
formula injection in downstream spreadsheet applications.

Encode direction: QuerySet → Serializer → Formatter → File/String
Decode direction: File/String → Formatter → Serializer → Models
"""

import csv
import json
from abc import ABC, abstractmethod
from io import BytesIO, StringIO
from typing import Any, Dict, List, Union

from openpyxl import Workbook, load_workbook


# Module imports
from plane.utils.csv_utils import sanitize_csv_row, sanitize_csv_value


class BaseFormatter(ABC):
    """Polymorphic interface for formatter implementations.

    Subclasses must override :meth:`encode`, :meth:`decode`, and the
    :attr:`extension` property so callers can treat JSON, CSV, and XLSX
    payloads uniformly. The contract is round-trip lossy on purpose:
    flattening, header prettification, and JSON-string fallbacks for nested
    values are owned by each subclass.
    """

    @abstractmethod
    def encode(self, data: List[Dict]) -> Union[str, bytes]:
        """Encode rows into the format's serialized string or byte payload."""
        pass

    @abstractmethod
    def decode(self, content: Union[str, bytes]) -> List[Dict]:
        """Decode a serialized payload back into a list of row dicts."""
        pass

    @property
    @abstractmethod
    def extension(self) -> str:
        """Return the canonical file extension for this format (without leading dot)."""
        pass


class JSONFormatter(BaseFormatter):
    """Encode and decode row data using the stdlib :mod:`json` module.

    Lists of dicts are serialized with :func:`json.dumps`, falling back to
    :class:`str` for non-JSON-native values (e.g. ``datetime``, ``UUID``) so
    DRF serializer output round-trips cleanly.
    """

    def __init__(self, indent: int = 2):
        """Initialize the JSON formatter with the indent width passed to :func:`json.dumps`."""
        self.indent = indent

    def encode(self, data: List[Dict]) -> str:
        """Serialize ``data`` as a JSON string with configured indentation."""
        return json.dumps(data, indent=self.indent, default=str)

    def decode(self, content: str) -> List[Dict]:
        """Parse a JSON string into the original list of row dicts."""
        return json.loads(content)

    @property
    def extension(self) -> str:
        """Return ``"json"`` so :class:`DataExporter` can compose ``<name>.json`` filenames."""
        return "json"


class CSVFormatter(BaseFormatter):
    """Encode and decode row data as CSV with optional nested-dict flattening.

    Flattens nested dicts using ``parent__child`` keys, JSON-encodes list
    values, prettifies headers (``created_by_name`` → ``Created By Name``)
    on encode, and reverses both transformations on decode. Every encoded
    value passes through :func:`plane.utils.csv_utils.sanitize_csv_value` to
    prevent CSV formula injection in downstream spreadsheet applications.
    """

    def __init__(self, flatten: bool = True, delimiter: str = ",", prettify_headers: bool = True):
        """Initialize the CSV formatter with flattening, delimiter, and header style options.

        Args:
            flatten: Whether to flatten nested dicts.
            delimiter: CSV delimiter character.
            prettify_headers: If True, transforms 'created_by_name' → 'Created By Name'.
        """
        self.flatten = flatten
        self.delimiter = delimiter
        self.prettify_headers = prettify_headers

    def _prettify_header(self, header: str) -> str:
        """Transform 'created_by_name' → 'Created By Name'."""
        return header.replace("_", " ").title()

    def _normalize_header(self, header: str) -> str:
        """Transform 'Display Name' → 'display_name' (reverse of prettify)."""
        return header.strip().lower().replace(" ", "_")

    def _flatten(self, row: Dict, parent_key: str = "") -> Dict:
        items = {}
        for key, value in row.items():
            new_key = f"{parent_key}__{key}" if parent_key else key
            if isinstance(value, dict):
                items.update(self._flatten(value, new_key))
            elif isinstance(value, list):
                items[new_key] = json.dumps(value)
            else:
                items[new_key] = value
        return items

    def _unflatten(self, row: Dict) -> Dict:
        result = {}
        for key, value in row.items():
            parts = key.split("__")
            current = result
            for part in parts[:-1]:
                current = current.setdefault(part, {})

            if isinstance(value, str):
                try:
                    parsed = json.loads(value)
                    if isinstance(parsed, (list, dict)):
                        value = parsed
                except (json.JSONDecodeError, TypeError):
                    pass

            current[parts[-1]] = value
        return result

    def encode(self, data: List[Dict]) -> str:
        """Serialize ``data`` as a CSV string with sanitized values and configured header style."""
        if not data:
            return ""

        if self.flatten:
            data = [self._flatten(row) for row in data]

        # Collect all unique field names in order
        fieldnames = []
        for row in data:
            for key in row.keys():
                if key not in fieldnames:
                    fieldnames.append(key)

        output = StringIO()

        if self.prettify_headers:
            # Create header mapping: original_key → Pretty Header
            header_map = {key: self._prettify_header(key) for key in fieldnames}
            pretty_headers = [header_map[key] for key in fieldnames]

            # Write pretty headers manually, then write data rows
            writer = csv.writer(output, delimiter=self.delimiter)
            writer.writerow(pretty_headers)

            # Write data rows in the same field order
            for row in data:
                writer.writerow(sanitize_csv_row([row.get(key, "") for key in fieldnames]))
        else:
            writer = csv.DictWriter(output, fieldnames=fieldnames, delimiter=self.delimiter)
            writer.writeheader()
            for row in data:
                writer.writerow({k: sanitize_csv_value(row.get(k, "")) for k in fieldnames})

        return output.getvalue()

    def decode(self, content: str, normalize_headers: bool = True) -> List[Dict]:
        """Decode CSV content into a list of row dicts.

        Args:
            content: CSV string.
            normalize_headers: If True, converts 'Display Name' → 'display_name'.
        """
        rows = list(csv.DictReader(StringIO(content), delimiter=self.delimiter))

        # Normalize headers: 'Email' → 'email', 'Display Name' → 'display_name'
        if normalize_headers:
            rows = [{self._normalize_header(k): v for k, v in row.items()} for row in rows]

        if self.flatten:
            rows = [self._unflatten(row) for row in rows]

        return rows

    @property
    def extension(self) -> str:
        """Return ``"csv"`` so :class:`DataExporter` can compose ``<name>.csv`` filenames."""
        return "csv"


class XLSXFormatter(BaseFormatter):
    """Encode and decode row data as XLSX workbooks using :mod:`openpyxl`.

    Writes and reads workbooks via in-memory :class:`io.BytesIO` buffers so
    no temporary files are required, joins list values with ``list_joiner``
    on encode, JSON-stringifies dict values, and best-effort restores nested
    structures on decode via JSON parsing.
    """

    def __init__(self, prettify_headers: bool = True, list_joiner: str = ", "):
        """Initialize the XLSX formatter with header style and list-joining options.

        Args:
            prettify_headers: If True, transforms 'created_by_name' → 'Created By Name'.
            list_joiner: String to join list values (default: ", ").
        """
        self.prettify_headers = prettify_headers
        self.list_joiner = list_joiner

    def _prettify_header(self, header: str) -> str:
        """Transform 'created_by_name' → 'Created By Name'."""
        return header.replace("_", " ").title()

    def _normalize_header(self, header: str) -> str:
        """Transform 'Display Name' → 'display_name' (reverse of prettify)."""
        return header.strip().lower().replace(" ", "_")

    def _format_value(self, value: Any) -> Any:
        """Format a value for XLSX cell."""
        if value is None:
            return ""
        if isinstance(value, list):
            return self.list_joiner.join(str(v) for v in value)
        if isinstance(value, dict):
            return json.dumps(value)
        return value

    def encode(self, data: List[Dict]) -> bytes:
        """Encode data to XLSX bytes."""
        wb = Workbook()
        ws = wb.active

        if not data:
            # Return empty workbook
            output = BytesIO()
            wb.save(output)
            output.seek(0)
            return output.getvalue()

        # Collect all unique field names in order
        fieldnames = []
        for row in data:
            for key in row.keys():
                if key not in fieldnames:
                    fieldnames.append(key)

        # Write header row
        if self.prettify_headers:
            headers = [self._prettify_header(key) for key in fieldnames]
        else:
            headers = fieldnames
        ws.append(headers)

        # Write data rows
        for row in data:
            ws.append([self._format_value(row.get(key, "")) for key in fieldnames])

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    def decode(self, content: bytes, normalize_headers: bool = True) -> List[Dict]:
        """Decode XLSX workbook bytes into a list of row dicts.

        Args:
            content: XLSX file bytes.
            normalize_headers: If True, converts 'Display Name' → 'display_name'.
        """
        wb = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
        ws = wb.active

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []

        # First row is headers
        headers = list(rows[0])
        if normalize_headers:
            headers = [self._normalize_header(str(h)) if h else "" for h in headers]

        # Convert remaining rows to dicts
        result = []
        for row in rows[1:]:
            row_dict = {}
            for i, value in enumerate(row):
                if i < len(headers) and headers[i]:
                    # Try to parse JSON strings back to lists/dicts
                    if isinstance(value, str):
                        try:
                            parsed = json.loads(value)
                            if isinstance(parsed, (list, dict)):
                                value = parsed
                        except (json.JSONDecodeError, TypeError):
                            pass
                    row_dict[headers[i]] = value
            result.append(row_dict)

        return result

    @property
    def extension(self) -> str:
        """Return ``"xlsx"`` so :class:`DataExporter` can compose ``<name>.xlsx`` filenames."""
        return "xlsx"
